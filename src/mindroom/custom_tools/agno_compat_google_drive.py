"""Agno Google Drive Office text extraction that keeps tables."""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

import agno.tools.google.drive as agno_google_drive

if TYPE_CHECKING:
    from collections.abc import Iterable, Iterator

    from pptx.shapes.base import BaseShape


def _extract_docx_text_with_tables(content_bytes: bytes) -> str:
    import docx  # noqa: PLC0415
    from docx.table import Table  # noqa: PLC0415

    lines: list[str] = []
    for block in docx.Document(io.BytesIO(content_bytes)).iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                cells = [cell.text for cell in row.cells]
                if any(cells):
                    lines.append("\t".join(cells))
        else:
            lines.append(block.text)
    return "\n".join(lines)


def _pptx_shape_lines(shapes: Iterable[BaseShape]) -> Iterator[str]:
    from pptx.shapes.graphfrm import GraphicFrame  # noqa: PLC0415
    from pptx.shapes.group import GroupShape  # noqa: PLC0415

    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _pptx_shape_lines(shape.shapes)
        elif isinstance(shape, GraphicFrame) and shape.has_table:
            for row in shape.table.rows:
                cells = [cell.text for cell in row.cells]
                if any(cells):
                    yield "\t".join(cells)
        elif shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                # Soft line breaks (Shift+Enter) come back as vertical tabs.
                text = paragraph.text.replace("\v", "\n")
                if text.strip():
                    yield text


def _extract_pptx_text_with_tables(content_bytes: bytes) -> str:
    from pptx import Presentation  # noqa: PLC0415

    lines: list[str] = []
    for number, slide in enumerate(Presentation(io.BytesIO(content_bytes)).slides, 1):
        lines.append(f"=== Slide {number} ===")
        lines.extend(_pptx_shape_lines(slide.shapes))
    return "\n".join(lines)


def install_office_table_extraction() -> None:
    """Replace Agno's Drive Office extractors with ones that keep tables; safe to call repeatedly."""
    # AGNO_COMPAT: Drive .docx text extraction drops tables.
    # Reason: Agno 3.0.9 `_extract_docx_text` reads only `document.paragraphs`, which excludes tables,
    # so `read_file` returns a document's text without any of its table cells.
    # Upstream issue: Tracking gap; no matching issue identified.
    # Upstream PR: https://github.com/agno-agi/agno/pull/10501, open.
    # Remove when: The pinned Agno `_extract_docx_text` returns table rows in document order.
    # Coverage: tests/test_google_drive_oauth_tool.py::test_google_drive_read_extracts_office_document_text.
    agno_google_drive._extract_docx_text = _extract_docx_text_with_tables  # ty: ignore[invalid-assignment]
    # AGNO_COMPAT: Drive .pptx text extraction drops tables, grouped shapes, and soft line breaks.
    # Reason: Agno 3.0.9 `_extract_pptx_text` reads only top-level shapes with a text frame and joins
    # paragraph runs, so slide tables and grouped shape text are missing from `read_file`, and values
    # separated by Shift+Enter run together.
    # Upstream issue: Tracking gap; no matching issue identified.
    # Upstream PR: None identified.
    # Remove when: The pinned Agno `_extract_pptx_text` returns table rows, grouped shape text, and
    # paragraph text with its line breaks.
    # Coverage: tests/test_google_drive_oauth_tool.py::test_google_drive_read_extracts_presentation_tables_and_groups.
    agno_google_drive._extract_pptx_text = _extract_pptx_text_with_tables  # ty: ignore[invalid-assignment]
